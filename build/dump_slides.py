import sys
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
want = set(int(x) for x in sys.argv[2].split(','))
W,H = p.slide_width, p.slide_height
for i, s in enumerate(p.slides):
    if i not in want: continue
    print(f"===== slide {i}")
    def walk(shapes, d=0):
        for sh in shapes:
            if sh.shape_type == 6: walk(sh.shapes, d+1); continue
            tag = f"{'  '*d}[{sh.shape_id} {sh.name} t={sh.shape_type} top={sh.top/H if sh.top is not None else 0:.2f} left={sh.left/W if sh.left is not None else 0:.2f}]"
            if sh.has_text_frame:
                print(tag, ' | '.join(repr(pp.text) for pp in sh.text_frame.paragraphs))
            elif getattr(sh,'has_table',False) and sh.has_table:
                print(tag, 'TABLE')
                for r in sh.table.rows: print('     ', [c.text for c in r.cells])
            else: print(tag)
    walk(s.shapes)
