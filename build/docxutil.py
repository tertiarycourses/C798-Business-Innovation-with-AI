import copy, re
from docx.oxml.ns import qn

def set_par(p, text):
    """Write text into the first non-empty run (or first run), blank the rest."""
    runs = p.runs
    if not runs:
        p.add_run(text); return
    tgt = next((r for r in runs if r.text), runs[0])
    for r in runs:
        r.text = text if r is tgt else ''

def set_cell(cell, text):
    ps = cell.paragraphs
    set_par(ps[0], text)
    for p in ps[1:]:
        p._p.getparent().remove(p._p)

def del_par(p):
    p._p.getparent().remove(p._p)

def del_row(table, row):
    table._tbl.remove(row._tr)

def drop_columns(table, idxs):
    """Remove columns idxs (0-based) and rescale remaining widths to the original total."""
    tbl = table._tbl
    grid = tbl.tblGrid
    cols = grid.findall(qn('w:gridCol'))
    total = sum(int(c.get(qn('w:w'))) for c in cols)
    for i in sorted(idxs, reverse=True):
        grid.remove(cols[i])
        for tr in tbl.findall(qn('w:tr')):
            tcs = tr.findall(qn('w:tc'))
            tr.remove(tcs[i])
    cols = grid.findall(qn('w:gridCol'))
    rem = sum(int(c.get(qn('w:w'))) for c in cols)
    new = [int(int(c.get(qn('w:w'))) * total / rem) for c in cols]
    for c, w in zip(cols, new): c.set(qn('w:w'), str(w))
    for tr in tbl.findall(qn('w:tr')):
        for tc, w in zip(tr.findall(qn('w:tc')), new):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is not None and tcPr.find(qn('w:tcW')) is not None:
                tcPr.find(qn('w:tcW')).set(qn('w:w'), str(w))
                tcPr.find(qn('w:tcW')).set(qn('w:type'), 'dxa')

def replace_everywhere(doc, subs):
    """Literal run-level replacement in body, tables, headers/footers."""
    def paras():
        yield from doc.paragraphs
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells: yield from c.paragraphs
        for s in doc.sections:
            for part in (s.header, s.footer, s.first_page_header, s.first_page_footer, s.even_page_header, s.even_page_footer):
                yield from part.paragraphs
                for t in part.tables:
                    for r in t.rows:
                        for c in r.cells: yield from c.paragraphs
    n = 0
    for p in paras():
        full = p.text
        for a, b in subs:
            if a in full:
                # try run-level first
                hit = False
                for r in p.runs:
                    if a in r.text:
                        r.text = r.text.replace(a, b); hit = True; n += 1
                if not hit:
                    set_par(p, full.replace(a, b)); n += 1
                full = p.text
    return n
