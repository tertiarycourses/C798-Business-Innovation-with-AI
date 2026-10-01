"""Post-convert fixes for the C798 deck (run after convert_deck.py)."""
import re, sys
from pptx import Presentation
path = sys.argv[1]
p = Presentation(path)
H, W = p.slide_height, p.slide_width

def set_paras(sh, lines):
    """Write one string per existing paragraph via runs[0]; blank trailing runs."""
    paras = sh.text_frame.paragraphs
    assert len(paras) >= len(lines), (sh.shape_id, len(paras), lines)
    for para, txt in zip(paras, lines):
        runs = para.runs
        runs[0].text = txt
        for r in runs[1:]: r.text = ''
    for para in paras[len(lines):]:
        for r in para.runs: r.text = ''

def shapes(s):
    def walk(ss):
        for sh in ss:
            if sh.shape_type == 6: yield from walk(sh.shapes)
            else: yield sh
    return list(walk(s.shapes))

def by_id(s, sid):
    return next(sh for sh in shapes(s) if sh.shape_id == sid)

slides = list(p.slides)
# 1. cover code line
set_paras(by_id(slides[0], 12), ['Course Code: C798', 'Conducted by Tertiary Infotech Academy Pte Ltd  ·  UEN 201200696W'])
# 2. schedule slide: assessment card -> wrap-up
set_paras(by_id(slides[6], 17), ['Wrap-up'])
set_paras(by_id(slides[6], 18), ['•  Course recap and Q&A', '•  Implementation commitments', '•  Course feedback · 17:30 close'])
# 3. feedback slide
fb = slides[159]
set_paras(by_id(fb, 4), ['COURSE FEEDBACK'])
set_paras(by_id(fb, 5), ['Course Feedback'])
set_paras(by_id(fb, 11), ['Share feedback', 'Complete the course feedback form provided by your trainer.'])
set_paras(by_id(fb, 16), ['What worked', 'Tell us which topics and labs helped you most.'])
set_paras(by_id(fb, 21), ['What to improve', 'Suggest one change that would make the course more useful.'])
set_paras(by_id(fb, 26), ['Certificate', 'Your certificate of completion will be emailed to you after the course.'])
# 4. strip A#/K# competency codes from eyebrows
code = r'[AK]\d(?:-[AK]\d)?'
for i, s in enumerate(slides):
    for sh in shapes(s):
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if re.fullmatch(rf'{code} · {code}', t):
            topic = re.search(r'Key Concepts — Topic (\d)', ' '.join(x.text_frame.text for x in shapes(s) if x.has_text_frame))
            set_paras(sh, [f'TOPIC 0{topic.group(1)} · KEY CONCEPTS' if topic else 'WRAP-UP'])
        elif re.search(rf' · {code} · {code} · ', t):
            set_paras(sh, [re.sub(rf' · {code} · {code}', '', t)])
# 5. footer page numbers follow the source convention (number = slide index)
for i, s in enumerate(slides):
    for sh in shapes(s):
        if sh.has_text_frame and sh.top and sh.top / H > 0.85 and sh.left / W > 0.5 and re.fullmatch(r'\d+', sh.text_frame.text.strip()):
            set_paras(sh, [str(i)])
p.save(path)
print('fixed', path)
