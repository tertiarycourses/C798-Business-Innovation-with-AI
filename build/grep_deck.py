import sys,re
from pptx import Presentation
p=Presentation(sys.argv[1]); rx=re.compile(sys.argv[2], re.I)
def texts(shapes):
    for sh in shapes:
        if sh.shape_type==6: yield from texts(sh.shapes); continue
        if sh.has_text_frame: yield sh, sh.text_frame.text
        if getattr(sh,'has_table',False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells: yield sh, c.text
for i,s in enumerate(p.slides):
    for sh,t in texts(s.shapes):
        if rx.search(t): print(i, 'SHAPE', sh.shape_id, repr(t[:160]))
    if s.has_notes_slide:
        t=s.notes_slide.notes_text_frame.text
        for m in rx.finditer(t): print(i,'NOTES',repr(t[max(0,m.start()-60):m.end()+60]))
